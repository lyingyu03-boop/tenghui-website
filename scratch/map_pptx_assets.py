import os
import pptx
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

pptx_path = r"c:\SYNC\TENG\騰煇_素材\ATeng\tenghui-website\小U-WEB\專欄文章\簡報檔-小磨坊斗六廠辦更新工程.pptx"
out_dir = r"c:\SYNC\TENG\騰煇_素材\ATeng\tenghui-website\img\projects\xiaomofang_assets"
os.makedirs(out_dir, exist_ok=True)

prs = Presentation(pptx_path)

def extract_pictures_from_shapes(shapes, slide_num, counter):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            counter[0] += 1
            image = shape.image
            ext = image.ext
            filename = f"slide_{slide_num:02d}_pic_{counter[0]:02d}.{ext}"
            filepath = os.path.join(out_dir, filename)
            with open(filepath, "wb") as f:
                f.write(image.blob)
            print(f"Slide {slide_num:02d} Pic {counter[0]}: Saved {filename} ({len(image.blob)} bytes)")
        elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            extract_pictures_from_shapes(shape.shapes, slide_num, counter)

print(f"Processing {len(prs.slides)} slides in PPTX...")
for i, slide in enumerate(prs.slides):
    slide_num = i + 1
    counter = [0]
    extract_pictures_from_shapes(slide.shapes, slide_num, counter)

print("Slide-by-slide image extraction finished!")
